from flask import Blueprint, render_template, request, jsonify, session, redirect, url_for, flash, send_file
from bson import ObjectId
from datetime import datetime, timedelta
from functools import wraps
import os
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from models.admin import Admin
from models.booking import Booking
from models.slot import Slot
from models.settings import Settings
from models.location import Location
from utils.slot_utils import sort_slots, slot_label, slot_start_datetime, rate_type_for_date, get_bangladesh_now
from db import mongo
from werkzeug.utils import secure_filename

admin_bp = Blueprint('admin', __name__)

def admin_required(f):
    """Decorator to require admin authentication"""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get('admin_logged_in'):
            return redirect(url_for('admin.login'))
        return f(*args, **kwargs)
    return decorated_function

def superadmin_required(f):
    """Decorator to require superadmin authentication"""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get('admin_logged_in'):
            return redirect(url_for('admin.login'))
        if session.get('admin_role') != 'superadmin':
            flash('Access denied: Super Administrator privileges required.', 'error')
            return redirect(url_for('admin.dashboard'))
        return f(*args, **kwargs)
    return decorated_function

def get_admin_location_filter(custom_location_id=None):
    """
    Returns a query dictionary filtering by location.
    If the logged-in admin is a branch admin, their location is strictly enforced.
    If custom_location_id is given (and admin is superadmin), filters by that location.
    """
    target_loc_id = session.get('admin_location_id') or (custom_location_id if custom_location_id and custom_location_id != 'all' else None)
    if not target_loc_id:
        return {}

    target_loc_str = str(target_loc_id)
    loc_obj_id = ObjectId(target_loc_str) if len(target_loc_str) == 24 else target_loc_str
    
    locs = Location.find_all()
    is_default_location = bool(locs and str(locs[0]['_id']) == target_loc_str)
    
    if is_default_location:
        return {
            '$or': [
                {'location_id': loc_obj_id},
                {'location_id': target_loc_str},
                {'location_id': {'$exists': False}},
                {'location_id': None}
            ]
        }
    else:
        return {
            '$or': [
                {'location_id': loc_obj_id},
                {'location_id': target_loc_str}
            ]
        }

def export_bookings_to_excel(bookings_list, locations_map, slots_map):
    """Generate Excel file in-memory for bookings"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bookings"
    headers = [
        'Booking ID', 'Date', 'Location', 'Time Slot', 'Name', 'Phone', 
        'Team Name', 'Payment Method', 'Amount', 'Status', 
        'Transaction ID', 'Notes', 'Created At'
    ]
    ws.append(headers)
    
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        
    for b in bookings_list:
        slot = slots_map.get(str(b.get('time_slot'))) if b.get('time_slot') else None
        slot_time = slot_label(slot)
        loc_id = b.get('location_id') or (slot.get('location_id') if slot else None)
        loc = locations_map.get(str(loc_id)) if loc_id else None
        loc_name = loc.get('name', 'N/A') if loc else 'N/A'
        
        row = [
            b.get('booking_id') or str(b.get('_id')),
            b.get('date', ''),
            loc_name,
            slot_time,
            b.get('name', ''),
            b.get('phone', ''),
            b.get('team_name', ''),
            b.get('payment_method', ''),
            b.get('amount', 0),
            b.get('status', ''),
            b.get('transaction_id', ''),
            b.get('notes', ''),
            str(b.get('created_at', ''))
        ]
        ws.append(row)
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

# ==================== Authentication ====================

@admin_bp.route('/login', methods=['GET', 'POST'])
def login():
    """Admin login"""
    if request.method == 'POST':
        username = request.form.get('username', '').strip().lower()
        password = request.form.get('password', '')
        
        admin = Admin.find_by_username(username)
        # If admin not found, attempt to initialize defaults and retry
        if not admin:
            try:
                Admin.init_defaults()
                admin = Admin.find_by_username(username)
            except Exception:
                admin = None
                
        if admin and Admin.verify_password(admin, password):
            session['admin_logged_in'] = True
            session['admin_username'] = admin.get('username', username)
            session['admin_role'] = admin.get('role', 'branch_admin' if admin.get('location_id') else 'superadmin')
            session['admin_location_id'] = str(admin['location_id']) if admin.get('location_id') else None
            
            if admin.get('location_id'):
                loc = Location.find_by_id(str(admin['location_id']))
                session['admin_location_name'] = loc.get('name') if loc else None
            else:
                session['admin_location_name'] = None
                
            return redirect(url_for('admin.dashboard'))
        else:
            flash('Invalid username or password', 'error')
    
    if session.get('admin_logged_in'):
        return redirect(url_for('admin.dashboard'))
    return render_template('admin/login.html')

@admin_bp.route('/logout')
def logout():
    """Admin logout"""
    session.pop('admin_logged_in', None)
    session.pop('admin_username', None)
    session.pop('admin_role', None)
    session.pop('admin_location_id', None)
    session.pop('admin_location_name', None)
    return redirect(url_for('admin.login'))

# ==================== Dashboard ====================

@admin_bp.route('/')
@admin_required
def dashboard():
    """Admin dashboard (scoped by branch for branch admins)"""
    now_bangladesh = get_bangladesh_now()
    today = now_bangladesh.strftime('%Y-%m-%d')
    loc_filter = get_admin_location_filter()
    
    # Today's bookings
    today_query = {'date': today}
    if loc_filter:
        today_query = {'$and': [{'date': today}, loc_filter]}
    today_bookings = Booking.count(today_query)
    
    # Upcoming bookings
    upcoming_base = {
        'date': {'$gte': today},
        'status': {'$in': ['pending', 'approved', 'paid']}
    }
    if loc_filter:
        upcoming_query = {'$and': [upcoming_base, loc_filter]}
    else:
        upcoming_query = upcoming_base
    upcoming_bookings = Booking.count(upcoming_query)
    
    # Calculate total earnings
    paid_base = {'status': 'paid'}
    if loc_filter:
        paid_query = {'$and': [paid_base, loc_filter]}
    else:
        paid_query = paid_base
    paid_bookings = Booking.find_all(paid_query)
    total_earnings = sum(b.get('amount', 0) for b in paid_bookings)
    
    # Match query for aggregations
    agg_match = {'status': {'$in': ['pending', 'approved', 'paid']}}
    if loc_filter:
        agg_match = {'$and': [{'status': {'$in': ['pending', 'approved', 'paid']}}, loc_filter]}
        
    # Popular slots
    pipeline = [
        {'$match': agg_match},
        {'$group': {'_id': '$time_slot', 'count': {'$sum': 1}}},
        {'$sort': {'count': -1}},
        {'$limit': 5}
    ]
    popular_slots = list(Booking.aggregate(pipeline))
    
    # Recent bookings
    recent_query = loc_filter if loc_filter else {}
    recent_bookings = Booking.find_all(recent_query, sort=[('created_at', -1)], limit=10)
    
    # Slot performance counts
    all_slot_counts = list(Booking.aggregate([
        {'$match': agg_match},
        {'$group': {'_id': '$time_slot', 'count': {'$sum': 1}}}
    ]))

    # Collect all slot IDs needed
    slot_ids = set()
    for slot_data in popular_slots:
        if slot_data.get('_id'):
            slot_ids.add(slot_data['_id'])
    for booking in recent_bookings:
        if booking.get('time_slot'):
            slot_ids.add(booking['time_slot'])
    for item in all_slot_counts:
        if item.get('_id'):
            slot_ids.add(item['_id'])
            
    slot_object_ids = []
    for sid in slot_ids:
        if isinstance(sid, ObjectId):
            slot_object_ids.append(sid)
        elif isinstance(sid, str) and len(sid) == 24:
            slot_object_ids.append(ObjectId(sid))
        else:
            slot_object_ids.append(sid)
            
    slots_list = Slot.find_all({'_id': {'$in': slot_object_ids}})
    slots_map = {str(s['_id']): s for s in slots_list}
    
    # Collect all location IDs
    loc_ids = set()
    for slot in slots_list:
        if slot.get('location_id'):
            loc_ids.add(slot['location_id'])
    for booking in recent_bookings:
        if booking.get('location_id'):
            loc_ids.add(booking['location_id'])
            
    loc_object_ids = []
    for lid in loc_ids:
        if isinstance(lid, ObjectId):
            loc_object_ids.append(lid)
        elif isinstance(lid, str) and len(lid) == 24:
            loc_object_ids.append(ObjectId(lid))
        else:
            loc_object_ids.append(lid)
            
    locations_list = Location.find_all({'_id': {'$in': loc_object_ids}})
    locations_map = {str(l['_id']): l for l in locations_list}
    
    def get_slot(sid):
        return slots_map.get(str(sid)) if sid else None
        
    def get_location(lid):
        return locations_map.get(str(lid)) if lid else None
    
    # Enrich popular slots
    for slot_data in popular_slots:
        slot = get_slot(slot_data['_id'])
        if slot:
            slot_data['time'] = slot_label(slot)
            slot_data['sport'] = slot.get('sport', 'N/A').title()
            loc_id = slot.get('location_id')
            loc = get_location(loc_id)
            slot_data['location_name'] = loc.get('name', 'N/A') if loc else 'N/A'
        else:
            slot_data['time'] = 'N/A'
            slot_data['sport'] = 'N/A'
            slot_data['location_name'] = 'N/A'
    
    # Enrich recent bookings
    for booking in recent_bookings:
        slot_id = booking.get('time_slot')
        slot = get_slot(slot_id)
        booking['slot_time'] = slot_label(slot)
        loc_id = booking.get('location_id') or (slot.get('location_id') if slot else None)
        loc = get_location(loc_id)
        booking['location_name'] = loc.get('name', 'N/A') if loc else 'N/A'
    
    performance_labels = []
    performance_values = []
    
    distribution = {
        'Morning': 0,    # 06:00 - 12:00
        'Afternoon': 0,  # 12:00 - 18:00
        'Evening': 0,    # 18:00 - 21:00
        'Night': 0       # 21:00 - 06:00
    }
    
    for item in all_slot_counts:
        slot = get_slot(item['_id'])
        if slot:
            time_str = slot.get('time', 'N/A')
            sport_name = slot.get('sport', 'football').title()
            performance_labels.append(f"{slot_label(slot)} ({sport_name})")
            performance_values.append(item['count'])
            
            try:
                start_hour = int(time_str.split(':')[0])
                if 6 <= start_hour < 12:
                    distribution['Morning'] += item['count']
                elif 12 <= start_hour < 18:
                    distribution['Afternoon'] += item['count']
                elif 18 <= start_hour < 21:
                    distribution['Evening'] += item['count']
                else:
                    distribution['Night'] += item['count']
            except (ValueError, IndexError):
                pass

    return render_template('admin/dashboard.html',
                         today_bookings=today_bookings,
                         upcoming_bookings=upcoming_bookings,
                         total_earnings=total_earnings,
                         popular_slots=popular_slots,
                         recent_bookings=recent_bookings,
                         performance_labels=performance_labels,
                         performance_values=performance_values,
                         distribution_labels=list(distribution.keys()),
                         distribution_values=list(distribution.values()))

# ==================== Bookings Management ====================

@admin_bp.route('/bookings')
@admin_required
def bookings():
    """View bookings (scoped by branch for branch admins)"""
    filter_type = request.args.get('filter', 'all')
    now_bangladesh = get_bangladesh_now()
    today = now_bangladesh.strftime('%Y-%m-%d')
    
    is_branch_admin = bool(session.get('admin_location_id'))
    if is_branch_admin:
        selected_location_id = session.get('admin_location_id')
        current_loc = Location.find_by_id(selected_location_id)
        locations_list = [current_loc] if current_loc else []
    else:
        selected_location_id = request.args.get('location_id', 'all')
        locations_list = Location.find_all()
        
    query_parts = []
    
    # Status / Date filter
    if filter_type == 'today':
        query_parts.append({'date': today})
    elif filter_type == 'pending':
        query_parts.append({'status': 'pending'})
    elif filter_type == 'approved':
        query_parts.append({'status': 'approved'})
        
    # Location filter
    loc_filter = get_admin_location_filter(selected_location_id)
    if loc_filter:
        query_parts.append(loc_filter)
        
    if len(query_parts) == 1:
        query = query_parts[0]
    elif len(query_parts) > 1:
        query = {'$and': query_parts}
    else:
        query = {}
    
    bookings_list = Booking.find_all(query, sort=[('created_at', -1)])
    
    # Collect all slot IDs
    slot_ids = set()
    for booking in bookings_list:
        if booking.get('time_slot'):
            slot_ids.add(booking['time_slot'])
            
    slot_object_ids = []
    for sid in slot_ids:
        if isinstance(sid, ObjectId):
            slot_object_ids.append(sid)
        elif isinstance(sid, str) and len(sid) == 24:
            slot_object_ids.append(ObjectId(sid))
        else:
            slot_object_ids.append(sid)
            
    slots_list = Slot.find_all({'_id': {'$in': slot_object_ids}})
    slots_map = {str(s['_id']): s for s in slots_list}
    
    # Collect all location IDs
    loc_ids = set()
    for slot in slots_list:
        if slot.get('location_id'):
            loc_ids.add(slot['location_id'])
    for booking in bookings_list:
        if booking.get('location_id'):
            loc_ids.add(booking['location_id'])
            
    loc_object_ids = []
    for lid in loc_ids:
        if isinstance(lid, ObjectId):
            loc_object_ids.append(lid)
        elif isinstance(lid, str) and len(lid) == 24:
            loc_object_ids.append(ObjectId(lid))
        else:
            loc_object_ids.append(lid)
            
    all_locations = Location.find_all({'_id': {'$in': loc_object_ids}})
    locations_map = {str(l['_id']): l for l in all_locations}
    
    def get_slot(sid):
        return slots_map.get(str(sid)) if sid else None
        
    def get_location(lid):
        return locations_map.get(str(lid)) if lid else None
        
    for booking in bookings_list:
        slot_id = booking.get('time_slot')
        slot = get_slot(slot_id)
        booking['slot_time'] = slot_label(slot)
        
        loc_id = booking.get('location_id') or (slot.get('location_id') if slot else None)
        loc = get_location(loc_id)
        booking['location_name'] = loc.get('name', 'N/A') if loc else 'N/A'
        
        if not booking.get('booking_id'):
            booking['booking_id'] = f"LEGACY_{str(booking['_id'])[-6:]}"
            
        try:
            date_str = booking.get('date')
            slot_time = slot['time'] if slot else None
            status = booking.get('status')
            played = False
            if date_str and slot_time:
                slot_start = slot_start_datetime(date_str, slot_time)
                played = now_bangladesh >= slot_start and status != 'cancelled'
            booking['status_display'] = 'played' if played else status
        except Exception:
            booking['status_display'] = booking.get('status')
    
    return render_template('admin/bookings.html',
                           bookings=bookings_list,
                           filter_type=filter_type,
                           locations=locations_list,
                           selected_location_id=selected_location_id,
                           is_branch_admin=is_branch_admin)

@admin_bp.route('/booking/<booking_id>/<action>', methods=['POST'])
@admin_required
def update_booking(booking_id, action):
    """Update booking status (verified by branch ownership for branch admin)"""
    booking = Booking.find_by_id(booking_id)
    if not booking:
        return jsonify({'error': 'Booking not found'}), 404
    
    # Verify branch admin authorization
    admin_loc_id = session.get('admin_location_id')
    if admin_loc_id:
        booking_loc_id = str(booking.get('location_id')) if booking.get('location_id') else None
        locs = Location.find_all()
        is_default = bool(locs and str(locs[0]['_id']) == str(admin_loc_id))
        
        if booking_loc_id != str(admin_loc_id) and not (is_default and not booking_loc_id):
            return jsonify({'error': 'Unauthorized: You can only update bookings for your assigned turf branch.'}), 403
    
    target_status = None
    if action == 'approve':
        target_status = 'approved'
    elif action == 'cancel':
        target_status = 'cancelled'
    elif action == 'paid':
        target_status = 'paid'
    elif action == 'pending':
        target_status = 'pending'
    
    if target_status and target_status in ['approved', 'paid', 'pending']:
        existing = Booking.find_by_date_and_slot(
            booking.get('date'),
            booking.get('time_slot'),
            exclude_booking_id=booking_id
        )
        if existing:
            return jsonify({'error': 'Cannot change status: This slot is already booked for this date by another active booking.'}), 400

    if target_status:
        Booking.update(booking_id, {'status': target_status})
    
    return jsonify({'success': True})

@admin_bp.route('/booking/create', methods=['GET', 'POST'])
@admin_required
def create_manual_booking():
    """Create manual booking (restricted to assigned branch for branch admins)"""
    admin_loc_id = session.get('admin_location_id')
    is_branch_admin = bool(admin_loc_id)
    
    if request.method == 'POST':
        data = request.form
        location_id = admin_loc_id if is_branch_admin else data.get('location_id')
        slot = Slot.find_by_id(data.get('time_slot'))
        
        if not slot:
            flash('Invalid time slot', 'error')
            return redirect(url_for('admin.create_manual_booking'))
        
        # Verify slot belongs to assigned location
        if is_branch_admin:
            slot_loc_id = str(slot.get('location_id')) if slot.get('location_id') else None
            locs = Location.find_all()
            is_default = bool(locs and str(locs[0]['_id']) == str(admin_loc_id))
            if slot_loc_id != str(admin_loc_id) and not (is_default and not slot_loc_id):
                flash('You can only create bookings for your assigned branch.', 'error')
                return redirect(url_for('admin.create_manual_booking'))
        
        # Check if slot is already booked for the selected date
        existing = Booking.find_by_date_and_slot(data.get('date'), data.get('time_slot'))
        if existing:
            flash('This slot is already booked for the selected date.', 'error')
            return redirect(url_for('admin.create_manual_booking'))
        
        # Ensure rate type matches date (weekend vs weekday)
        try:
            date_obj = datetime.strptime(data.get('date'), '%Y-%m-%d')
            expected_rate = rate_type_for_date(date_obj)
            if slot.get('rate_type', 'weekday') != expected_rate:
                loc_id_match = slot.get('location_id') or (ObjectId(location_id) if location_id and len(location_id) == 24 else None)
                matching = Slot.get_collection().find_one({
                    'location_id': loc_id_match,
                    'sport': slot.get('sport'),
                    'time': slot.get('time'),
                    'venue': slot.get('venue'),
                    'rate_type': expected_rate,
                    'enabled': True
                })
                if matching:
                    slot = matching
        except Exception as e:
            print(f"Rate type matching warning: {e}")

        booking_data = {
            'date': data.get('date'),
            'time_slot': str(slot['_id']),
            'location_id': str(slot.get('location_id')) if slot.get('location_id') else location_id,
            'name': data.get('name'),
            'phone': data.get('phone'),
            'team_name': data.get('team_name', ''),
            'payment_method': data.get('payment_method'),
            'transaction_id': data.get('transaction_id', ''),
            'amount': slot['price'] * (1 - float(slot.get('discount_percent', 0)) / 100) if slot.get('discount_active') else slot['price'],
            'status': 'approved',
            'notes': data.get('notes', '')
        }
        
        Booking.create(booking_data)
        flash('Booking created successfully', 'success')
        return redirect(url_for('admin.bookings'))
    
    if is_branch_admin:
        assigned_loc = Location.find_by_id(admin_loc_id)
        locations = [assigned_loc] if assigned_loc else []
        slots_query = {'enabled': True, 'location_id': ObjectId(admin_loc_id) if len(admin_loc_id) == 24 else admin_loc_id}
        slots = Slot.find_all(slots_query)
        # If no slots under exact ObjectId, check without location if default
        locs = Location.find_all()
        if not slots and locs and str(locs[0]['_id']) == str(admin_loc_id):
            slots = Slot.find_all({'enabled': True})
    else:
        locations = Location.find_all()
        slots = Slot.find_all({'enabled': True})
        
    slots = sort_slots(slots)
    return render_template('admin/create_booking.html', slots=slots, locations=locations, is_branch_admin=is_branch_admin, current_location_id=admin_loc_id)

# ==================== Slots Management ====================

@admin_bp.route('/slots')
@admin_required
def slots():
    """View slots (scoped by branch for branch admins)"""
    admin_loc_id = session.get('admin_location_id')
    is_branch_admin = bool(admin_loc_id)
    
    if is_branch_admin:
        loc_obj = ObjectId(admin_loc_id) if len(admin_loc_id) == 24 else admin_loc_id
        football_slots = Slot.find_all({'sport': 'football', 'location_id': loc_obj})
        badminton_slots = Slot.find_all({'sport': 'badminton', 'location_id': loc_obj})
        current_loc = Location.find_by_id(admin_loc_id)
        locations = [current_loc] if current_loc else []
    else:
        football_slots = Slot.find_all({'sport': 'football'})
        badminton_slots = Slot.find_all({'sport': 'badminton'})
        locations = Location.find_all()
    
    football_slots = sort_slots(football_slots)
    badminton_slots = sort_slots(badminton_slots)
    
    return render_template('admin/slots.html',
                           football_slots=football_slots,
                           badminton_slots=badminton_slots,
                           locations=locations,
                           is_branch_admin=is_branch_admin,
                           current_location_id=admin_loc_id)

@admin_bp.route('/slot/create', methods=['POST'])
@admin_required
def create_slot():
    """Create new slot"""
    data = request.json
    admin_loc_id = session.get('admin_location_id')
    if admin_loc_id:
        data['location_id'] = admin_loc_id
        
    Slot.create(data)
    return jsonify({'success': True})

@admin_bp.route('/slot/<slot_id>/update', methods=['POST'])
@admin_required
def update_slot(slot_id):
    """Update slot"""
    slot = Slot.find_by_id(slot_id)
    if not slot:
        return jsonify({'error': 'Slot not found'}), 404
        
    admin_loc_id = session.get('admin_location_id')
    if admin_loc_id:
        if str(slot.get('location_id')) != str(admin_loc_id):
            return jsonify({'error': 'Unauthorized: Slot belongs to another branch.'}), 403
            
    data = request.json
    update_dict = {
        'time': data.get('time'),
        'price': float(data.get('price')),
        'sport': data.get('sport', 'football'),
        'rate_type': data.get('rate_type', 'weekday'),
        'enabled': data.get('enabled', True),
        'discount_active': data.get('discount_active', False),
        'discount_percent': float(data.get('discount_percent', 0.0)),
        'venue': data.get('venue', '')
    }
    if not admin_loc_id and 'location_id' in data:
        update_dict['location_id'] = data['location_id']
        
    Slot.update(slot_id, update_dict)
    return jsonify({'success': True})

# ==================== Calendar ====================

@admin_bp.route('/calendar')
@admin_required
def calendar():
    """Calendar view (passes branch id for branch admins)"""
    admin_loc_id = session.get('admin_location_id')
    return render_template('admin/calendar.html', current_location_id=admin_loc_id or '')

# ==================== Gallery Management ====================

@admin_bp.route('/gallery')
@admin_required
def gallery():
    """Gallery management"""
    try:
        images = list(mongo.db.gallery.find({'type': 'image'}))
        videos = list(mongo.db.gallery.find({'type': 'video'}))
    except Exception:
        images = []
        videos = []
    return render_template('admin/gallery.html', images=images, videos=videos)

@admin_bp.route('/gallery/upload', methods=['POST'])
@admin_required
def upload_gallery():
    """Upload gallery item"""
    if 'file' not in request.files:
        return jsonify({'error': 'No file uploaded'}), 400
    
    file = request.files['file']
    file_type = request.form.get('type', 'image')
    
    if file.filename == '':
        return jsonify({'error': 'No file selected'}), 400
    
    if os.environ.get('CLOUDINARY_URL'):
        import cloudinary
        import cloudinary.uploader
        try:
            res = cloudinary.uploader.upload(file, folder="gallery", resource_type="auto")
            url = res.get('secure_url')
            public_id = res.get('public_id')
            
            mongo.db.gallery.insert_one({
                'filename': file.filename,
                'path': url,
                'cloudinary_public_id': public_id,
                'type': file_type,
                'uploaded_at': get_bangladesh_now()
            })
            return jsonify({'success': True, 'filename': file.filename})
        except Exception as e:
            return jsonify({'error': f'Cloudinary upload error: {str(e)}'}), 500
    else:
        filename = secure_filename(file.filename)
        filepath = os.path.join('static/gallery', filename)
        file.save(filepath)
        
        mongo.db.gallery.insert_one({
            'filename': filename,
            'path': f'gallery/{filename}',
            'type': file_type,
            'uploaded_at': get_bangladesh_now()
        })
        return jsonify({'success': True, 'filename': filename})

@admin_bp.route('/gallery/delete/<gallery_id>', methods=['POST'])
@admin_required
def delete_gallery_item(gallery_id):
    """Delete gallery item"""
    item = mongo.db.gallery.find_one({'_id': ObjectId(gallery_id)})
    if item:
        if item.get('cloudinary_public_id'):
            import cloudinary
            import cloudinary.uploader
            try:
                cloudinary.uploader.destroy(item['cloudinary_public_id'], resource_type=item.get('type', 'image'))
            except Exception as e:
                print(f"Error deleting from Cloudinary: {e}")
        else:
            filepath = os.path.join('static', item['path'])
            if os.path.exists(filepath):
                os.remove(filepath)
        mongo.db.gallery.delete_one({'_id': ObjectId(gallery_id)})
    return jsonify({'success': True})

# ==================== Payments Management ====================

@admin_bp.route('/payments')
@admin_required
def payments():
    """Payments management (scoped by branch for branch admins)"""
    loc_filter = get_admin_location_filter()
    query = loc_filter if loc_filter else {}
    bookings_list = Booking.find_all(query, sort=[('created_at', -1)])
    
    slot_ids = set()
    for booking in bookings_list:
        if booking.get('time_slot'):
            slot_ids.add(booking['time_slot'])
            
    slot_object_ids = []
    for sid in slot_ids:
        if isinstance(sid, ObjectId):
            slot_object_ids.append(sid)
        elif isinstance(sid, str) and len(sid) == 24:
            slot_object_ids.append(ObjectId(sid))
        else:
            slot_object_ids.append(sid)
            
    slots_list = Slot.find_all({'_id': {'$in': slot_object_ids}})
    slots_map = {str(s['_id']): s for s in slots_list}
    
    for booking in bookings_list:
        slot_id = booking.get('time_slot')
        slot = slots_map.get(str(slot_id)) if slot_id else None
        booking['slot_time'] = slot_label(slot)
    return render_template('admin/payments.html', bookings=bookings_list)

# ==================== Carousel Management ====================

@admin_bp.route('/settings/carousel/upload', methods=['POST'])
@superadmin_required
def upload_carousel_image():
    """Upload carousel image (superadmin only)"""
    if 'file' not in request.files:
        return jsonify({'error': 'No file uploaded'}), 400
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No file selected'}), 400
    
    if os.environ.get('CLOUDINARY_URL'):
        import cloudinary
        import cloudinary.uploader
        try:
            res = cloudinary.uploader.upload(file, folder="carousel")
            url = res.get('secure_url')
            
            settings = Settings.find_one()
            images = settings.get('carousel_images', []) if settings else []
            images.append(url)
            Settings.update({'carousel_images': images})
            
            return jsonify({'success': True, 'filename': url})
        except Exception as e:
            return jsonify({'error': f'Cloudinary upload error: {str(e)}'}), 500
    else:
        filename = secure_filename(f"{int(get_bangladesh_now().timestamp())}_{file.filename}")
        filepath = os.path.join('static/carousel', filename)
        file.save(filepath)
        
        settings = Settings.find_one()
        images = settings.get('carousel_images', []) if settings else []
        images.append(filename)
        Settings.update({'carousel_images': images})
        
        return jsonify({'success': True, 'filename': filename})

@admin_bp.route('/settings/carousel/delete/<path:filename>', methods=['POST'])
@superadmin_required
def delete_carousel_image(filename):
    """Delete carousel image (superadmin only)"""
    if filename.startswith('http://') or filename.startswith('https://'):
        import cloudinary
        import cloudinary.uploader
        try:
            if 'image/upload/' in filename:
                parts = filename.split('image/upload/')[-1].split('/')
                if parts[0].startswith('v') and parts[0][1:].isdigit():
                    public_id_with_ext = '/'.join(parts[1:])
                else:
                    public_id_with_ext = '/'.join(parts)
                public_id = public_id_with_ext.rsplit('.', 1)[0]
                cloudinary.uploader.destroy(public_id)
            elif 'video/upload/' in filename:
                parts = filename.split('video/upload/')[-1].split('/')
                if parts[0].startswith('v') and parts[0][1:].isdigit():
                    public_id_with_ext = '/'.join(parts[1:])
                else:
                    public_id_with_ext = '/'.join(parts)
                public_id = public_id_with_ext.rsplit('.', 1)[0]
                cloudinary.uploader.destroy(public_id, resource_type="video")
        except Exception as e:
            print(f"Error deleting carousel image from Cloudinary: {e}")
    else:
        filepath = os.path.join('static/carousel', secure_filename(filename))
        if os.path.exists(filepath):
            os.remove(filepath)
    
    settings = Settings.find_one()
    images = settings.get('carousel_images', []) if settings else []
    if filename in images:
        images.remove(filename)
        Settings.update({'carousel_images': images})
        
    return jsonify({'success': True})

# ==================== Settings ====================

@admin_bp.route('/settings', methods=['GET', 'POST'])
@superadmin_required
def settings():
    """Settings management (superadmin only)"""
    if request.method == 'POST':
        Settings.update({
            'turf_name': request.form.get('turf_name'),
            'address': request.form.get('address'),
            'phone': request.form.get('phone'),
            'whatsapp': request.form.get('whatsapp'),
            'email': request.form.get('email'),
            'facebook': request.form.get('facebook'),
            'rules': request.form.get('rules'),
            'homepage_content': request.form.get('homepage_content')
        })
        flash('Settings updated successfully', 'success')
        return redirect(url_for('admin.settings'))
    
    settings = Settings.find_one()
    return render_template('admin/settings.html', settings=settings)

# ==================== Export ====================

@admin_bp.route('/export')
@superadmin_required
def export():
    """Export bookings to Excel (superadmin only)"""
    admin_loc_id = session.get('admin_location_id')
    loc_filter = get_admin_location_filter()
    
    # Fetch all relevant bookings
    query = loc_filter if loc_filter else {}
    bookings_list = Booking.find_all(query, sort=[('date', -1)])
    
    slots_list = Slot.find_all()
    slots_map = {str(s['_id']): s for s in slots_list}
    
    locations_list = Location.find_all()
    locations_map = {str(l['_id']): l for l in locations_list}
    
    excel_stream = export_bookings_to_excel(bookings_list, locations_map, slots_map)
    
    branch_tag = f"_{session.get('admin_location_name', 'branch').replace(' ', '_').lower()}" if admin_loc_id else ""
    filename = f"bookings{branch_tag}.xlsx"
    
    return send_file(
        excel_stream,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

# ==================== Messages Management ====================

@admin_bp.route('/messages')
@admin_required
def messages():
    """View contact messages"""
    from models.contact_message import ContactMessage
    messages = ContactMessage.find_all()
    for msg in messages:
        if 'created_at' in msg and isinstance(msg['created_at'], datetime):
            msg['display_date'] = msg['created_at'].strftime('%Y-%m-%d %H:%M')
        else:
            msg['display_date'] = 'N/A'
    return render_template('admin/messages.html', messages=messages)

@admin_bp.route('/messages/delete/<message_id>', methods=['POST'])
@admin_required
def delete_message(message_id):
    """Delete a contact message"""
    from models.contact_message import ContactMessage
    ContactMessage.delete(message_id)
    flash('Message deleted successfully.', 'success')
    return redirect(url_for('admin.messages'))

# ==================== Locations Management ====================

@admin_bp.route('/locations')
@admin_required
def locations():
    """View locations (superadmin or branch-scoped)"""
    admin_loc_id = session.get('admin_location_id')
    if admin_loc_id:
        loc = Location.find_by_id(admin_loc_id)
        locations_list = [loc] if loc else []
    else:
        locations_list = Location.find_all()
    return render_template('admin/locations.html', locations=locations_list)

@admin_bp.route('/location/create', methods=['GET', 'POST'])
@superadmin_required
def create_location():
    """Create a new location (superadmin only)"""
    if request.method == 'POST':
        image_url = request.form.get('image_url', '')
        if 'image_file' in request.files:
            file = request.files['image_file']
            if file and file.filename != '':
                if os.environ.get('CLOUDINARY_URL'):
                    import cloudinary
                    import cloudinary.uploader
                    try:
                        res = cloudinary.uploader.upload(file, folder="locations")
                        image_url = res.get('secure_url')
                    except Exception as e:
                        flash(f'Cloudinary upload error: {str(e)}', 'error')
                else:
                    filename = secure_filename(f"{int(get_bangladesh_now().timestamp())}_{file.filename}")
                    os.makedirs('static/uploads', exist_ok=True)
                    filepath = os.path.join('static/uploads', filename)
                    file.save(filepath)
                    image_url = f"uploads/{filename}"

        data = {
            'name': request.form.get('name'),
            'address': request.form.get('address'),
            'phone': request.form.get('phone'),
            'whatsapp': request.form.get('whatsapp'),
            'email': request.form.get('email'),
            'rules': request.form.get('rules'),
            'homepage_content': request.form.get('homepage_content'),
            'enabled': request.form.get('enabled') == 'true',
            'online_booking_enabled': request.form.get('online_booking_enabled') == 'true',
            'image_url': image_url,
            'facebook_url': request.form.get('facebook_url', '')
        }
        Location.create(data)
        flash('Location created successfully', 'success')
        return redirect(url_for('admin.locations'))
    
    return render_template('admin/create_location.html')

@admin_bp.route('/location/<location_id>/update', methods=['GET', 'POST'])
@admin_required
def update_location(location_id):
    """Update a location"""
    admin_loc_id = session.get('admin_location_id')
    if admin_loc_id and str(admin_loc_id) != str(location_id):
        flash('Access denied: You can only edit your assigned turf branch.', 'error')
        return redirect(url_for('admin.locations'))

    location = Location.find_by_id(location_id)
    if not location:
        flash('Location not found', 'error')
        return redirect(url_for('admin.locations'))
        
    if request.method == 'POST':
        image_url = request.form.get('image_url', '')
        if 'image_file' in request.files:
            file = request.files['image_file']
            if file and file.filename != '':
                if os.environ.get('CLOUDINARY_URL'):
                    import cloudinary
                    import cloudinary.uploader
                    try:
                        res = cloudinary.uploader.upload(file, folder="locations")
                        image_url = res.get('secure_url')
                    except Exception as e:
                        flash(f'Cloudinary upload error: {str(e)}', 'error')
                else:
                    filename = secure_filename(f"{int(get_bangladesh_now().timestamp())}_{file.filename}")
                    os.makedirs('static/uploads', exist_ok=True)
                    filepath = os.path.join('static/uploads', filename)
                    file.save(filepath)
                    image_url = f"uploads/{filename}"

        data = {
            'name': request.form.get('name'),
            'address': request.form.get('address'),
            'phone': request.form.get('phone'),
            'whatsapp': request.form.get('whatsapp'),
            'email': request.form.get('email'),
            'rules': request.form.get('rules'),
            'homepage_content': request.form.get('homepage_content'),
            'enabled': request.form.get('enabled') == 'true',
            'online_booking_enabled': request.form.get('online_booking_enabled') == 'true',
            'image_url': image_url,
            'facebook_url': request.form.get('facebook_url', '')
        }
        Location.update(location_id, data)
        flash('Location updated successfully', 'success')
        return redirect(url_for('admin.locations'))
        
    return render_template('admin/edit_location.html', location=location)

@admin_bp.route('/location/<location_id>/delete', methods=['POST'])
@superadmin_required
def delete_location(location_id):
    """Delete a location (superadmin only)"""
    Location.delete(location_id)
    flash('Location deleted successfully', 'success')
    return redirect(url_for('admin.locations'))

# ==================== Admin Users Management (Superadmin only) ====================

@admin_bp.route('/users')
@superadmin_required
def admin_users():
    """List all admin accounts"""
    admins = Admin.find_all()
    locations = Location.find_all()
    locations_map = {str(l['_id']): l.get('name', 'N/A') for l in locations}
    
    for a in admins:
        lid = a.get('location_id')
        a['location_name'] = locations_map.get(str(lid), 'All Locations (Superadmin)' if not lid else 'Unknown')
        
    return render_template('admin/admin_users.html', admins=admins, locations=locations)

@admin_bp.route('/users/create', methods=['POST'])
@superadmin_required
def create_admin_user():
    """Create a new admin user"""
    username = request.form.get('username', '').strip().lower()
    password = request.form.get('password', '').strip()
    full_name = request.form.get('full_name', '').strip()
    role = request.form.get('role', 'branch_admin')
    location_id = request.form.get('location_id', '') if role == 'branch_admin' else None
    
    if not username or not password:
        flash('Username and password are required.', 'error')
        return redirect(url_for('admin.admin_users'))
        
    if Admin.find_by_username(username):
        flash(f"An admin with username '{username}' already exists.", 'error')
        return redirect(url_for('admin.admin_users'))
        
    Admin.create(
        username=username,
        password=password,
        location_id=location_id if location_id else None,
        role=role,
        full_name=full_name or username
    )
    flash(f"Admin '{username}' created successfully.", 'success')
    return redirect(url_for('admin.admin_users'))

@admin_bp.route('/users/<admin_id>/update', methods=['POST'])
@superadmin_required
def update_admin_user(admin_id):
    """Update an existing admin user"""
    admin = Admin.find_by_id(admin_id)
    if not admin:
        flash('Admin user not found.', 'error')
        return redirect(url_for('admin.admin_users'))
        
    full_name = request.form.get('full_name', '').strip()
    role = request.form.get('role', 'branch_admin')
    location_id = request.form.get('location_id', '') if role == 'branch_admin' else None
    new_password = request.form.get('new_password', '').strip()
    
    update_data = {
        'full_name': full_name or admin.get('username'),
        'role': role,
        'location_id': location_id if location_id else None
    }
    if new_password:
        update_data['password'] = new_password
        
    Admin.update(admin_id, update_data)
    flash(f"Admin '{admin.get('username')}' updated successfully.", 'success')
    return redirect(url_for('admin.admin_users'))

@admin_bp.route('/users/<admin_id>/delete', methods=['POST'])
@superadmin_required
def delete_admin_user(admin_id):
    """Delete an admin user"""
    admin = Admin.find_by_id(admin_id)
    if not admin:
        flash('Admin user not found.', 'error')
        return redirect(url_for('admin.admin_users'))
        
    if admin.get('username') == session.get('admin_username'):
        flash('You cannot delete your own logged-in admin account.', 'error')
        return redirect(url_for('admin.admin_users'))
        
    Admin.delete(admin_id)
    flash(f"Admin '{admin.get('username')}' deleted successfully.", 'success')
    return redirect(url_for('admin.admin_users'))
